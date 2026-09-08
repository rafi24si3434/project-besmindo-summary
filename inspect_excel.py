import openpyxl
import xlrd
import mysql.connector

# Cek isi lembar file Excel dulu
wb_npt = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\System\NPT SEPTEMBER  2026 SYS.xlsx', data_only=True)
print("Sheet NPT:", wb_npt.sheetnames)

wb_dr = xlrd.open_workbook(r'C:\Users\LENOVO\Downloads\System\Daily Report SEPTEMBER 2026 SYS.xls')
print("Sheet Daily Report:", wb_dr.sheet_names())

# Cek isi salah satu sheet NPT (BMS 01)
ws = wb_npt['BMS 01']
print("\nBMS 01 NPT Header Row 2:", [ws.cell(2, c).value for c in range(1, 15)])
print("BMS 01 NPT Header Row 3:", [ws.cell(3, c).value for c in range(1, 15)])
print("BMS 01 NPT Header Row 4:", [ws.cell(4, c).value for c in range(1, 15)])
for r in range(5, 12):
    row_v = [ws.cell(r, c).value for c in range(1, 10)]
    if any(row_v):
        print(f"BMS 01 Row {r}:", row_v)

# Cek isi salah satu sheet Daily Report (BMS 01)
ws_dr = wb_dr.sheet_by_name('BMS 01')
for r in range(15):
    rv = ws_dr.row_values(r)
    if any(rv):
        print(f"DR Row {r}:", rv[:8])
